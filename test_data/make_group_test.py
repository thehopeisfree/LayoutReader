"""Generate a test PPTX with various group configurations.

Scenarios:
  1. Simple group: two rectangles grouped together
  2. Rotated group: group with 30-degree rotation
  3. Nested group: group inside a group
  4. Scaled group: chExt != ext (child coordinate space differs)
  5. Standalone shapes (not grouped) as reference

Also renders a reference PNG with Pillow for verification.
"""
from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

OUT_DIR = Path(__file__).parent
SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
PNG_W = 1920
PNG_H = 1080
SX = PNG_W / SLIDE_W_EMU
SY = PNG_H / SLIDE_H_EMU


def emu2px(x_emu: int, y_emu: int) -> tuple[float, float]:
    return x_emu * SX, y_emu * SY


def add_group_with_two_rects(slide, prs, group_x, group_y, group_w, group_h,
                              rect1_color, rect2_color, group_name="Group"):
    """Add a group shape containing two rectangles side by side."""
    from pptx.oxml import parse_xml
    from lxml import etree

    nsmap = {
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    }

    half_w = group_w // 2

    # Build group XML manually
    grpSp_xml = f'''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:nvGrpSpPr>
        <p:cNvPr id="0" name="{group_name}"/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm>
          <a:off x="{group_x}" y="{group_y}"/>
          <a:ext cx="{group_w}" cy="{group_h}"/>
          <a:chOff x="0" y="0"/>
          <a:chExt cx="{group_w}" cy="{group_h}"/>
        </a:xfrm>
      </p:grpSpPr>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="GRect1"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="0" y="0"/>
            <a:ext cx="{half_w}" cy="{group_h}"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="{rect1_color}"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>Left</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="GRect2"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{half_w}" y="0"/>
            <a:ext cx="{half_w}" cy="{group_h}"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="{rect2_color}"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>Right</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:grpSp>'''

    grpSp_el = parse_xml(grpSp_xml)
    slide._element.find(qn('p:cSld')).find(qn('p:spTree')).append(grpSp_el)
    return grpSp_el


def add_rotated_group(slide, prs, group_x, group_y, group_w, group_h,
                       rot_deg, rect1_color, rect2_color, group_name="RotGroup"):
    """Add a rotated group."""
    from pptx.oxml import parse_xml

    half_w = group_w // 2
    rot_60000 = int(rot_deg * 60000)

    grpSp_xml = f'''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                              xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:nvGrpSpPr>
        <p:cNvPr id="0" name="{group_name}"/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm rot="{rot_60000}">
          <a:off x="{group_x}" y="{group_y}"/>
          <a:ext cx="{group_w}" cy="{group_h}"/>
          <a:chOff x="0" y="0"/>
          <a:chExt cx="{group_w}" cy="{group_h}"/>
        </a:xfrm>
      </p:grpSpPr>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="RotRect1"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="0" y="0"/>
            <a:ext cx="{half_w}" cy="{group_h}"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="{rect1_color}"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>RotL</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="RotRect2"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="{half_w}" y="0"/>
            <a:ext cx="{half_w}" cy="{group_h}"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="{rect2_color}"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>RotR</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:grpSp>'''

    grpSp_el = parse_xml(grpSp_xml)
    slide._element.find(qn('p:cSld')).find(qn('p:spTree')).append(grpSp_el)


def add_nested_group(slide, prs):
    """Add a nested group (group inside group)."""
    from pptx.oxml import parse_xml

    # Outer group: at (6000000, 3500000), size 5000000 x 3000000
    # Inner group: at (0, 0) in child coords, size 5000000 x 1500000
    #   - two rects inside inner group
    # Another rect directly in outer group at (0, 1500000)

    grpSp_xml = '''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                             xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                             xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:nvGrpSpPr>
        <p:cNvPr id="0" name="OuterGroup"/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm>
          <a:off x="6000000" y="3500000"/>
          <a:ext cx="5000000" cy="3000000"/>
          <a:chOff x="0" y="0"/>
          <a:chExt cx="5000000" cy="3000000"/>
        </a:xfrm>
      </p:grpSpPr>

      <!-- Inner group (top half) -->
      <p:grpSp>
        <p:nvGrpSpPr>
          <p:cNvPr id="0" name="InnerGroup"/>
          <p:cNvGrpSpPr/>
          <p:nvPr/>
        </p:nvGrpSpPr>
        <p:grpSpPr>
          <a:xfrm>
            <a:off x="0" y="0"/>
            <a:ext cx="5000000" cy="1500000"/>
            <a:chOff x="0" y="0"/>
            <a:chExt cx="5000000" cy="1500000"/>
          </a:xfrm>
        </p:grpSpPr>
        <p:sp>
          <p:nvSpPr>
            <p:cNvPr id="0" name="InnerRect1"/>
            <p:cNvSpPr/>
            <p:nvPr/>
          </p:nvSpPr>
          <p:spPr>
            <a:xfrm>
              <a:off x="0" y="0"/>
              <a:ext cx="2500000" cy="1500000"/>
            </a:xfrm>
            <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
            <a:solidFill><a:srgbClr val="FF6600"/></a:solidFill>
          </p:spPr>
          <p:txBody>
            <a:bodyPr/>
            <a:lstStyle/>
            <a:p><a:r><a:t>Inner1</a:t></a:r></a:p>
          </p:txBody>
        </p:sp>
        <p:sp>
          <p:nvSpPr>
            <p:cNvPr id="0" name="InnerRect2"/>
            <p:cNvSpPr/>
            <p:nvPr/>
          </p:nvSpPr>
          <p:spPr>
            <a:xfrm>
              <a:off x="2500000" y="0"/>
              <a:ext cx="2500000" cy="1500000"/>
            </a:xfrm>
            <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
            <a:solidFill><a:srgbClr val="0066FF"/></a:solidFill>
          </p:spPr>
          <p:txBody>
            <a:bodyPr/>
            <a:lstStyle/>
            <a:p><a:r><a:t>Inner2</a:t></a:r></a:p>
          </p:txBody>
        </p:sp>
      </p:grpSp>

      <!-- Direct child rect (bottom half) -->
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="OuterRect"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="0" y="1500000"/>
            <a:ext cx="5000000" cy="1500000"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="00CC66"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>OuterChild</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:grpSp>'''

    grpSp_el = parse_xml(grpSp_xml)
    slide._element.find(qn('p:cSld')).find(qn('p:spTree')).append(grpSp_el)


def add_scaled_group(slide, prs):
    """Add a group where chExt != ext (scaled child coordinate space)."""
    from pptx.oxml import parse_xml

    # Group at (500000, 3500000), ext=4000000x2500000
    # But chExt=8000000x5000000 (child space is 2x larger)
    # So children in child coords get scaled down by 0.5

    grpSp_xml = '''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
                             xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                             xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:nvGrpSpPr>
        <p:cNvPr id="0" name="ScaledGroup"/>
        <p:cNvGrpSpPr/>
        <p:nvPr/>
      </p:nvGrpSpPr>
      <p:grpSpPr>
        <a:xfrm>
          <a:off x="500000" y="3500000"/>
          <a:ext cx="4000000" cy="2500000"/>
          <a:chOff x="0" y="0"/>
          <a:chExt cx="8000000" cy="5000000"/>
        </a:xfrm>
      </p:grpSpPr>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="ScaledRect1"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="0" y="0"/>
            <a:ext cx="4000000" cy="5000000"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="CC44CC"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>Scaled1</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
      <p:sp>
        <p:nvSpPr>
          <p:cNvPr id="0" name="ScaledRect2"/>
          <p:cNvSpPr/>
          <p:nvPr/>
        </p:nvSpPr>
        <p:spPr>
          <a:xfrm>
            <a:off x="4000000" y="0"/>
            <a:ext cx="4000000" cy="5000000"/>
          </a:xfrm>
          <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
          <a:solidFill><a:srgbClr val="44CCCC"/></a:solidFill>
        </p:spPr>
        <p:txBody>
          <a:bodyPr/>
          <a:lstStyle/>
          <a:p><a:r><a:t>Scaled2</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:grpSp>'''

    grpSp_el = parse_xml(grpSp_xml)
    slide._element.find(qn('p:cSld')).find(qn('p:spTree')).append(grpSp_el)


def make_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # --- Reference: standalone shape (top-left) ---
    ref = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(500000), Emu(200000), Emu(2000000), Emu(800000),
    )
    ref.fill.solid()
    ref.fill.fore_color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    ref.text_frame.text = "Reference (no group)"

    # --- 1. Simple group (top-center) ---
    add_group_with_two_rects(
        slide, prs,
        group_x=3500000, group_y=200000,
        group_w=4000000, group_h=1200000,
        rect1_color="4488CC", rect2_color="CC4444",
        group_name="SimpleGroup",
    )

    # --- 2. Rotated group 30 deg (middle-left) ---
    add_rotated_group(
        slide, prs,
        group_x=500000, group_y=1800000,
        group_w=4000000, group_h=1200000,
        rot_deg=30,
        rect1_color="44CC44", rect2_color="CCCC44",
        group_name="RotatedGroup30",
    )

    # --- 3. Nested group (right side) ---
    add_nested_group(slide, prs)

    # --- 4. Scaled group (bottom-left) ---
    add_scaled_group(slide, prs)

    pptx_path = OUT_DIR / "test_group.pptx"
    prs.save(str(pptx_path))
    print(f"Saved {pptx_path}")
    return pptx_path


def make_reference_png() -> Path:
    """Draw approximate reference PNG."""
    img = Image.new("RGB", (PNG_W, PNG_H), (255, 255, 255))
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        font = ImageFont.load_default()

    # Reference rect
    x1, y1 = emu2px(500000, 200000)
    x2, y2 = emu2px(2500000, 1000000)
    draw.rectangle([x1, y1, x2, y2], fill=(0xAA, 0xAA, 0xAA))
    draw.text((x1+5, y1+5), "Reference", fill=(0,0,0), font=font)

    # Simple group: two rects side by side
    gx, gy = emu2px(3500000, 200000)
    gx2, gy2 = emu2px(5500000, 1400000)
    draw.rectangle([gx, gy, gx2, gy2], fill=(0x44, 0x88, 0xCC))
    gx3, _ = emu2px(5500000, 200000)
    gx4, _ = emu2px(7500000, 200000)
    draw.rectangle([gx3, gy, gx4, gy2], fill=(0xCC, 0x44, 0x44))

    # Scaled group: two halves
    sx1, sy1 = emu2px(500000, 3500000)
    sx2, sy2 = emu2px(2500000, 6000000)
    draw.rectangle([sx1, sy1, sx2, sy2], fill=(0xCC, 0x44, 0xCC))
    sx3, _ = emu2px(2500000, 3500000)
    sx4, _ = emu2px(4500000, 3500000)
    draw.rectangle([sx3, sy1, sx4, sy2], fill=(0x44, 0xCC, 0xCC))

    # Nested group
    nx1, ny1 = emu2px(6000000, 3500000)
    nx2, ny2 = emu2px(8500000, 5000000)
    draw.rectangle([nx1, ny1, nx2, ny2], fill=(0xFF, 0x66, 0x00))
    nx3, _ = emu2px(8500000, 3500000)
    nx4, _ = emu2px(11000000, 3500000)
    draw.rectangle([nx3, ny1, nx4, ny2], fill=(0x00, 0x66, 0xFF))
    nx5, ny3 = emu2px(6000000, 5000000)
    nx6, ny4 = emu2px(11000000, 6500000)
    draw.rectangle([nx5, ny3, nx6, ny4], fill=(0x00, 0xCC, 0x66))

    png_path = OUT_DIR / "test_group.png"
    img.save(str(png_path))
    print(f"Saved {png_path}")
    return png_path


if __name__ == "__main__":
    make_pptx()
    make_reference_png()
    print("Done. Run:")
    print("  python pptx_bbox.py test_data/test_group.pptx 1 test_data/test_group.png --out-json test_data/group_out.json --overlay test_data/group_overlay.png")
