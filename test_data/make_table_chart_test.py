"""Generate a test PPTX with table and chart elements.

Elements:
  1. Title textbox
  2. A 3x4 table
  3. A bar chart
  4. A pie chart
  5. A reference rectangle
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData

OUT_DIR = Path(__file__).parent
SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
PNG_W = 1920
PNG_H = 1080
SX = PNG_W / SLIDE_W_EMU
SY = PNG_H / SLIDE_H_EMU


def emu2px(x_emu: int, y_emu: int) -> tuple[float, float]:
    return x_emu * SX, y_emu * SY


def make_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # --- 1. Title textbox ---
    txBox = slide.shapes.add_textbox(Emu(500000), Emu(200000), Emu(11000000), Emu(600000))
    txBox.text_frame.text = "Table & Chart Test Slide"
    for para in txBox.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(28)
            run.font.color.rgb = RGBColor(0, 0, 0)

    # --- 2. Table (3 cols x 4 rows) ---
    rows, cols = 4, 3
    tbl = slide.shapes.add_table(
        rows, cols,
        Emu(500000), Emu(1000000),
        Emu(5000000), Emu(3000000),
    )
    table = tbl.table
    # Header row
    headers = ["Category", "Q1", "Q2"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
    # Data rows
    data = [
        ["Product A", "120", "150"],
        ["Product B", "80", "200"],
        ["Product C", "95", "110"],
    ]
    for r, row_data in enumerate(data, 1):
        for c, val in enumerate(row_data):
            table.cell(r, c).text = val

    # --- 3. Bar chart ---
    chart_data = CategoryChartData()
    chart_data.categories = ["Product A", "Product B", "Product C"]
    chart_data.add_series("Q1", (120, 80, 95))
    chart_data.add_series("Q2", (150, 200, 110))

    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Emu(6000000), Emu(1000000),
        Emu(5500000), Emu(3000000),
        chart_data,
    )

    # --- 4. Pie chart ---
    pie_data = CategoryChartData()
    pie_data.categories = ["Desktop", "Mobile", "Tablet"]
    pie_data.add_series("Share", (55, 35, 10))

    slide.shapes.add_chart(
        XL_CHART_TYPE.PIE,
        Emu(500000), Emu(4200000),
        Emu(5000000), Emu(2500000),
        pie_data,
    )

    # --- 5. Reference rectangle ---
    ref = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Emu(6000000), Emu(4200000),
        Emu(5500000), Emu(2500000),
    )
    ref.fill.solid()
    ref.fill.fore_color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    ref.text_frame.text = "Reference Shape"

    pptx_path = OUT_DIR / "test_table_chart.pptx"
    prs.save(str(pptx_path))
    print(f"Saved {pptx_path}")
    return pptx_path


def make_reference_png() -> Path:
    img = Image.new("RGB", (PNG_W, PNG_H), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        font = ImageFont.load_default()

    # Title
    x1, y1 = emu2px(500000, 200000)
    x2, y2 = emu2px(11500000, 800000)
    draw.rectangle([x1, y1, x2, y2], outline=(0, 0, 0))
    draw.text((x1+5, y1+5), "Table & Chart Test Slide", fill=(0, 0, 0), font=font)

    # Table area
    x1, y1 = emu2px(500000, 1000000)
    x2, y2 = emu2px(5500000, 4000000)
    draw.rectangle([x1, y1, x2, y2], fill=(230, 230, 250))
    draw.text((x1+5, y1+5), "[TABLE]", fill=(0, 0, 0), font=font)

    # Bar chart area
    x1, y1 = emu2px(6000000, 1000000)
    x2, y2 = emu2px(11500000, 4000000)
    draw.rectangle([x1, y1, x2, y2], fill=(250, 230, 230))
    draw.text((x1+5, y1+5), "[BAR CHART]", fill=(0, 0, 0), font=font)

    # Pie chart area
    x1, y1 = emu2px(500000, 4200000)
    x2, y2 = emu2px(5500000, 6700000)
    draw.rectangle([x1, y1, x2, y2], fill=(230, 250, 230))
    draw.text((x1+5, y1+5), "[PIE CHART]", fill=(0, 0, 0), font=font)

    # Reference
    x1, y1 = emu2px(6000000, 4200000)
    x2, y2 = emu2px(11500000, 6700000)
    draw.rectangle([x1, y1, x2, y2], fill=(0xDD, 0xDD, 0xDD))
    draw.text((x1+5, y1+5), "Reference Shape", fill=(0, 0, 0), font=font)

    png_path = OUT_DIR / "test_table_chart.png"
    img.save(str(png_path))
    print(f"Saved {png_path}")
    return png_path


if __name__ == "__main__":
    make_pptx()
    make_reference_png()
